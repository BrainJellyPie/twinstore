"""Spreadsheet errors and absent cached results cannot become real POS defaults."""

import io
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
import pytest

from store_emulator.application.pos import parse_pos_csv

HEADERS = ["receipt_id", "sku", "quantity", "unit_price", "hour", "date", "store_id", "register_id", "notes"]
VALUES = ["r1", "tea", 2, 250, 12, "2026-10-23", "shop", "till", ""]


def workbook(column, value, *, cached=None, sheet="first", nonstandard=False, header=False, row=2):
    book = Workbook()
    active = book.active
    active.append(HEADERS)
    active.append(VALUES)
    active.append(["r2", "coffee", 1, 300, 12, "2026-10-23", "shop", "till", ""])
    target = active if sheet == "first" else book.create_sheet("Unused")
    target.cell(1 if header else row, HEADERS.index(column) + 1, value)
    source = io.BytesIO()
    book.save(source)
    if cached is None and not nonstandard:
        return source.getvalue()
    output = io.BytesIO()
    with ZipFile(source) as archive, ZipFile(output, "w", ZIP_DEFLATED) as rewritten:
        for info in archive.infolist():
            content = archive.read(info)
            name = info.filename
            if name == "xl/worksheets/sheet1.xml":
                if cached is not None:
                    root = ET.fromstring(content)
                    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
                    address = f"{chr(65 + HEADERS.index(column))}{1 if header else row}"
                    cell = next(element for element in root.iter(f"{ns}c") if element.get("r") == address)
                    cell.set("t", cached[0])
                    value_node = cell.find(f"{ns}v")
                    if value_node is None:
                        value_node = ET.SubElement(cell, f"{ns}v")
                    value_node.text = cached[1]
                    content = ET.tostring(root)
                if nonstandard:
                    name = "xl/worksheets/custom.data"
            if nonstandard:
                content = content.replace(b"worksheets/sheet1.xml", b"worksheets/custom.data")
            rewritten.writestr(name, content)
    return output.getvalue()


@pytest.mark.parametrize("column", HEADERS[:-1])
@pytest.mark.parametrize("value", ['="tea"', "#REF!", "#DIV/0!"])
def test_unreadable_imported_cells_cannot_drop_receipts_or_default_values(column, value):
    with pytest.raises(ValueError, match="POS worksheet row 2, column"):
        parse_pos_csv(workbook(column, value), "pos.xlsx")


@pytest.mark.parametrize("value", ['="unit_price"', "#REF!"])
def test_unreadable_optional_headers_are_not_silently_omitted(value):
    with pytest.raises(ValueError, match="POS worksheet row 1, column 4"):
        parse_pos_csv(workbook("unit_price", value, header=True), "pos.xlsx")


@pytest.mark.parametrize("column,cached,expected", [
    ("sku", ("str", "00123"), ("00123", 2, 250)),
    ("quantity", ("n", "0"), ("tea", 0, 250)),
    ("unit_price", ("n", "125"), ("tea", 2, 125)),
    ("unit_price", ("str", ""), ("tea", 2, 0)),
])
def test_saved_formula_results_are_used_including_zero_and_empty_text(column, cached, expected):
    lines = parse_pos_csv(workbook(column, "=1+1", cached=cached), "pos.xlsx")
    assert (lines[0].sku, lines[0].quantity, lines[0].unit_price) == expected


def test_a_saved_empty_sku_formula_still_means_an_unused_row():
    lines = parse_pos_csv(workbook("sku", '=IF(1,"","")', cached=("str", "")), "pos.xlsx")
    assert [line.transaction_id for line in lines] == ["r2"]


@pytest.mark.parametrize("value", ["=1+1", "#REF!"])
@pytest.mark.parametrize("sheet,column", [("first", "notes"), ("second", "sku")])
def test_unused_notes_and_other_worksheets_do_not_block_valid_receipts(value, sheet, column):
    lines = parse_pos_csv(workbook(column, value, sheet=sheet), "pos.xlsx")
    assert [(line.transaction_id, line.sku) for line in lines] == [("r1", "tea"), ("r2", "coffee")]


def test_nonstandard_worksheet_part_is_checked_using_the_imported_sheet():
    with pytest.raises(ValueError, match="POS worksheet row 2, column 2"):
        parse_pos_csv(workbook("sku", '="tea"', nonstandard=True), "pos.xlsx")


@pytest.mark.parametrize("error", ["#SPILL!", "#CALC!", "#BLOCKED!"])
def test_newer_excel_errors_are_detected_by_cell_type(error):
    with pytest.raises(ValueError, match="POS worksheet row 2, column 2"):
        parse_pos_csv(workbook("sku", "=1+1", cached=("e", error)), "pos.xlsx")


def test_uncached_formula_beyond_the_readers_last_visible_row_is_not_dropped():
    with pytest.raises(ValueError, match="POS worksheet row 4, column 2"):
        parse_pos_csv(workbook("sku", '="tea"', row=4), "pos.xlsx")


def test_a_lower_priority_unused_identifier_alias_does_not_block_the_selected_sku():
    book = Workbook()
    sheet = book.active
    sheet.append(["receipt_id", "sku", "barcode"])
    sheet.append(["r", "tea", "#REF!"])
    output = io.BytesIO()
    book.save(output)
    lines = parse_pos_csv(output.getvalue(), "pos.xlsx")
    assert [(line.transaction_id, line.sku) for line in lines] == [("r", "tea")]
