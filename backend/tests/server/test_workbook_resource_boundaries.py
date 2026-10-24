"""Small ZIP uploads must not expand into unbounded pandas tables."""

import io
import json
import re
from unittest.mock import Mock
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from store_emulator.application import pos, supplemental, workbooks
from tests.server.test_project_templates import _meta, client


def workbook(ranges=((3, 2),)):
    book = Workbook()
    for index, (rows, columns) in enumerate(ranges):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = f"Data {index}"
        sheet.append(["receipt_id", "sku"])
        sheet.append(["01", "tea"])
        sheet.cell(rows, columns, "last")
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def transform(data, kind):
    output = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        if kind == "unrelated_broken_xml":
            target.writestr("customXml/broken.xml", "<broken")
        for info in source.infolist():
            content = source.read(info)
            name = info.filename
            if name == "xl/worksheets/sheet1.xml":
                if kind == "false_dimension":
                    content = re.sub(rb'<dimension ref="[^"]+"\s*/>', b'<dimension ref="A1:B3"/>', content)
                elif kind == "no_dimension":
                    content = re.sub(rb'<dimension ref="[^"]+"\s*/>', b'', content)
                elif kind == "nonstandard_part":
                    name = "xl/worksheets/custom.data"
                elif kind == "implicit_columns":
                    content = re.sub(rb'<c r="[^"]+"', b'<c', content)
            if kind == "nonstandard_part":
                content = content.replace(b"worksheets/sheet1.xml", b"worksheets/custom.data")
            target.writestr(name, content)
    return output.getvalue()


@pytest.mark.parametrize("kind", ["normal", "false_dimension", "no_dimension", "nonstandard_part", "unrelated_broken_xml"])
def test_sparse_workbook_is_rejected_before_dataframe_loading(monkeypatch, kind):
    data = transform(workbook(((50000, 2000),)), kind)
    assert len(data) < 10_000
    reader = Mock(side_effect=AssertionError("pandas must not allocate the sparse rectangle"))
    monkeypatch.setattr(pos.pd, "read_excel", reader)
    monkeypatch.setattr(pos.pd, "ExcelFile", reader)
    with pytest.raises(workbooks.WorkbookLimitError, match="empty gaps"):
        pos.parse_pos_csv(data, "visits.XLSX")
    assert supplemental._load_frames(data, ".xlsx") == {}
    reader.assert_not_called()


@pytest.mark.parametrize("filename", ["visits.xlsx", "visits.XLSX"])
def test_wizard_rejects_before_creating_project_or_starting_generation(client, filename):
    api, generation = client
    before = api.get("/api/projects").json()
    response = api.post("/api/projects/create-store", data={"meta": json.dumps(_meta())},
                        files={"supplemental": (filename, workbook(((50000, 2000),)),
                                               "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 413 and "empty gaps" in response.text
    generation.assert_not_called()
    assert api.get("/api/projects").json() == before


@pytest.mark.parametrize("ranges,allowed", [
    (((1000, 2000),), True), (((1001, 2000),), False),
    (((1000, 1000), (1000, 1000)), True), (((1000, 1000), (1000, 1001)), False),
])
def test_budget_counts_empty_gaps_and_all_sheets_at_exact_boundary(ranges, allowed):
    data = workbook(ranges)
    if allowed:
        workbooks.validate_xlsx_budget(data)
    else:
        with pytest.raises(workbooks.WorkbookLimitError):
            workbooks.validate_xlsx_budget(data)


def test_archive_expansion_is_rejected_before_parsing_any_xml(monkeypatch):
    data = workbook()
    monkeypatch.setattr(workbooks, "MAX_WORKBOOK_BYTES", 1000)
    parse = Mock(side_effect=AssertionError("archive budget must precede decompression"))
    monkeypatch.setattr(workbooks, "iterparse", parse)
    with pytest.raises(workbooks.WorkbookLimitError, match="expands beyond"):
        workbooks.validate_xlsx_budget(data)
    parse.assert_not_called()


@pytest.mark.parametrize("kind", ["normal", "false_dimension", "no_dimension", "nonstandard_part", "implicit_columns"])
def test_ordinary_workbooks_keep_all_rows_and_identifiers(kind):
    data = transform(workbook(), kind)
    workbooks.validate_xlsx_budget(data)
    frame = supplemental._load_frames(data, ".xlsx", preserve_identifiers=True)["Data 0"]
    assert frame.shape == (2, 2)
    assert frame.iloc[0].tolist() == ["01", "tea"]
    lines = pos.parse_pos_csv(data, "visits.xlsx")
    assert len(lines) == 1 and lines[0].transaction_id == "01" and lines[0].sku == "tea"


def test_invalid_workbook_keeps_existing_failure_paths():
    workbooks.validate_xlsx_budget(b"not an Excel workbook")
    assert supplemental._load_frames(b"not an Excel workbook", ".xlsx") == {}
